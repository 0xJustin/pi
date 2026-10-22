"""Mechanical checks on a saved deck: what can be measured without judgment. Looking at the render is still required.
usage: python -m slidestyle.check deck.pptx   (also run by Deck.save)"""
import itertools, re, sys
from pathlib import Path
from matplotlib import ft2font
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from . import layouts as L

SKILL = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(SKILL.parent / "figures"))
from figstyle.palette import allowed                     # noqa: E402
from figstyle.check import ID_PATTERNS                   # noqa: E402

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
FONTS = SKILL.parent / "figures" / "fonts"
_FT = {}
TITLES = (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE)
OUR_FONTS = {"Source Sans 3", "Source Sans 3 Semibold", "+mj-lt", "+mn-lt"}
MIN_PT, MAX_WORDS, DPI = 14, 40, 150


def _font(bold):
    f = FONTS / ("SourceSans3-Semibold.ttf" if bold else "SourceSans3-Regular.ttf")
    if f not in _FT:
        _FT[f] = ft2font.FT2Font(str(f))
    return _FT[f]


def _width_pt(text, pt, bold):
    ft = _font(bold)
    ft.set_size(pt, 72)
    ft.set_text(text)
    return ft.get_width_height()[0] / 64


def _lines(text, pt, width, bold):
    """Lines a paragraph wraps to in a box `width` points wide (greedy word wrap, the font's real advances)."""
    n, line = 1, ""
    for word in text.split():
        trial = f"{line} {word}".strip()
        if line and _width_pt(trial, pt, bold) > width:
            n, line = n + 1, word
        else:
            line = trial
    return n


def _lvl_size(el, level):
    """sz (in pt) set on a list style's level `level`, or None."""
    if el is None:
        return None
    lvl = el.find(f"{A}lvl{level + 1}pPr/{A}defRPr")
    return int(lvl.get("sz")) / 100 if lvl is not None and lvl.get("sz") else None


def _size(slide, shape, para, run):
    """Font size in points, following PowerPoint's inheritance: run, shape, layout placeholder, master text styles."""
    if run is not None and run.font.size:
        return run.font.size.pt
    lvl = para.level
    own = _lvl_size(shape._element.find(f".//{A}lstStyle"), lvl)
    if own:
        return own
    if shape.is_placeholder:
        idx = shape.placeholder_format.idx
        for lp in slide.slide_layout.placeholders:
            if lp.placeholder_format.idx == idx:
                got = _lvl_size(lp._element.find(f".//{A}lstStyle"), lvl)
                if got:
                    return got
        styles = slide.slide_layout.slide_master._element.find(f"{P}txStyles")
        which = "titleStyle" if shape.placeholder_format.type in TITLES else "bodyStyle"
        return _lvl_size(styles.find(f"{P}{which}"), lvl) or 18
    return 18


def _inset(shape, side, default):
    bp = shape._element.find(f".//{A}bodyPr")
    v = bp.get(side) if bp is not None else None
    if v is None and shape.is_placeholder:
        return 0   # the template's placeholders all have zero insets
    return int(v) / 12700 if v is not None else default


def _text_issues(slide, shape, where):
    issues = []
    tf = shape.text_frame
    is_title = shape.is_placeholder and shape.placeholder_format.type in TITLES
    width = shape.width / 12700 - _inset(shape, "lIns", 7.2) - _inset(shape, "rIns", 7.2)
    height = shape.height / 12700 - _inset(shape, "tIns", 3.6) - _inset(shape, "bIns", 3.6)
    used, lines_total, smallest = 0.0, 0, None
    for k, para in enumerate(tf.paragraphs):
        text = "".join(r.text for r in para.runs) or para.text
        pt = _size(slide, shape, para, para.runs[0] if para.runs else None)
        smallest = pt if smallest is None else min(smallest, pt)
        indent = 27 * para.level
        n = _lines(text, pt, width - indent, is_title) if text.strip() else 1
        lines_total += n
        used += n * pt * 1.326 + (0 if k == 0 else (12 if para.level == 0 else 4))
        for r in para.runs:
            face = r.font.name
            if face and face not in OUR_FONTS:
                issues.append(f"{where}: font {face!r} is not Source Sans 3")
            rgb = r._r.find(f"{A}rPr/{A}solidFill/{A}srgbClr")
            if rgb is not None and not allowed("#" + rgb.get("val")):
                issues.append(f"{where}: text color #{rgb.get('val')} is not a quantity hue or a gray")
    if used > height + 2:
        issues.append(f"{where}: text needs {used / 72:.2f} in but its box is {height / 72:.2f} in tall; cut words "
                      "or split the slide")
    if is_title:
        if lines_total > 2:
            issues.append(f"{where}: the claim wraps to {lines_total} lines; a claim fits in two")
        if tf.text.rstrip().endswith("."):
            issues.append(f"{where}: the claim ends with a period")
    if smallest is not None and smallest < MIN_PT:
        issues.append(f"{where}: text at {smallest:g} pt; nothing on a slide goes under {MIN_PT} pt")
    for pat, what in ID_PATTERNS:
        if re.search(pat, tf.text):
            issues.append(f"{where}: {tf.text[:40]!r} looks like {what}; name what it stands for")
            break
    return issues


def check(path):
    prs = Presentation(str(path))
    W, H = prs.slide_width, prs.slide_height
    issues = []
    for n, slide in enumerate(prs.slides, 1):
        layout = slide.slide_layout.name
        if layout not in L.LAYOUTS:
            issues.append(f"slide {n}: layout {layout!r} is not from the dark template")
        words, boxes = 0, []
        for shape in slide.shapes:
            name = shape.name
            where = f"slide {n} {name!r}"
            if shape.is_placeholder and shape.placeholder_format.type == PP_PLACEHOLDER.SLIDE_NUMBER:
                continue
            if shape.left < -12700 or shape.top < -12700 or shape.left + shape.width > W + 12700 \
                    or shape.top + shape.height > H + 12700:
                issues.append(f"{where} runs off the slide")
            if shape.is_placeholder and shape.placeholder_format.type == PP_PLACEHOLDER.PICTURE \
                    and shape.shape_type == MSO_SHAPE_TYPE.PLACEHOLDER and not hasattr(shape, "image"):
                issues.append(f"{where}: empty figure frame (PowerPoint shows its prompt)")
                continue
            if shape.has_text_frame:
                if not shape.text_frame.text.strip():
                    if shape.is_placeholder:
                        issues.append(f"{where}: empty placeholder (PowerPoint shows its prompt); delete it or fill it")
                    continue
                words += len(shape.text_frame.text.split())
                issues += _text_issues(slide, shape, where)
                if not shape.is_placeholder:
                    boxes.append(shape)
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE or hasattr(shape, "crop_left"):
                crop = max(abs(getattr(shape, f"crop_{s}", 0) or 0) for s in ("left", "right", "top", "bottom"))
                if crop > 0.005:
                    issues.append(f"{where}: figure cropped by {crop:.0%} to fit its frame; make it at the frame's "
                                  "size (the layout's fs.slide_figure call)")
                try:
                    px = shape.image.size[0]
                    if px / (shape.width / 914400) < DPI:
                        issues.append(f"{where}: figure is {px / (shape.width / 914400):.0f} dpi on the slide; "
                                      f"save at {DPI}+")
                except (AttributeError, ValueError):
                    pass
        for a, b in itertools.combinations(boxes, 2):
            if a.left < b.left + b.width and b.left < a.left + a.width and a.top < b.top + b.height \
                    and b.top < a.top + a.height:
                issues.append(f"slide {n}: text boxes {a.name!r} and {b.name!r} overlap")
        if layout != "Recap / asks" and words > MAX_WORDS:
            issues.append(f"slide {n} carries {words} words; one claim per slide, the rest goes in the notes")
    return issues


if __name__ == "__main__":
    found = check(sys.argv[1])
    for i in found:
        print("SLIDE CHECK:", i)
    print(f"{len(found)} issues")
