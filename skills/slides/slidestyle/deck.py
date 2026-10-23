"""Build a deck on the dark template: one method per layout, figures dropped into frames made for them.

    d = Deck()                                   # or Deck("weekly.pptx") to add to a running deck
    d.claim_figure("The MLP policy keeps walking after the connectome-constrained model falls", "walk.png")
    d.save("talk.pptx")                          # runs the slide checks and prints every issue
"""
import copy, sys
from pathlib import Path
from pptx import Presentation
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn
from . import layouts as L

SKILL = Path(__file__).resolve().parent.parent
TEMPLATE = SKILL / "template" / "dark.pptx"
FIGURE_IDX = {"figure": 13, "left": 13, "right": 14}


class Deck:
    def __init__(self, path=None):
        self.prs = Presentation(str(path or TEMPLATE))
        self.layouts = {l.name: l for l in self.prs.slide_layouts}
        missing = set(L.LAYOUTS) - set(self.layouts)
        if missing:
            raise ValueError(f"{path} is not built on the dark template (no layouts {sorted(missing)})")

    # one method per layout ------------------------------------------------------------------------------------

    def title(self, title, subtitle=None, notes=None):
        s = self._slide("Title", title, notes)
        if subtitle:
            s.placeholders[1].text_frame.text = subtitle
        return s

    def section(self, title, notes=None):
        return self._slide("Section", title, notes)

    def claim_figure(self, claim, figure, notes=None, poster=None):
        """`figure`: a PNG made with fs.slide_figure(1.0, 0.78), or an MP4 (its poster frame is the middle unless given)."""
        s = self._slide("Claim + figure", claim, notes)
        self._place(s, "figure", figure, poster)
        return s

    def two_figures(self, claim, left, right, notes=None):
        """Two PNGs or MP4s, each made with fs.slide_figure(0.5, 0.78): one model or condition per side."""
        s = self._slide("Two figures", claim, notes)
        self._place(s, "left", left)
        self._place(s, "right", right)
        return s

    def figure(self, figure, notes=None, poster=None):
        """The whole slide is the figure (fs.slide_figure(1.0, 1.0)): for mature talks where the visual speaks."""
        s = self._slide("Figure only", None, notes)
        self._place(s, "figure", figure, poster)
        return s

    def claim(self, claim, notes=None):
        """A claim and an empty slide, for anything composed by hand on top of it."""
        return self._slide("Claim only", claim, notes)

    def recap_asks(self, claim, recap, asks, notes=None):
        """Weekly-meeting recap: what changed since last time (left), and what you need from the room (right).
        Items are strings; a nested list holds sub-items of the item before it."""
        s = self._slide("Recap / asks", claim, notes)
        self._items(s.placeholders[1], recap, "Since last time")
        self._items(s.placeholders[2], asks, "Asks")
        return s

    # running decks --------------------------------------------------------------------------------------------

    def move(self, slides, to):
        """Move slides (in order) to position `to` (0 = first): this week's block goes right after the title."""
        ids = self.prs.slides._sldIdLst
        entries = {int(e.get("id")): e for e in ids}
        for k, s in enumerate(slides):
            e = entries[s.slide_id]
            ids.remove(e)
            ids.insert(to + k, e)

    def save(self, path, check=True):
        """Number the slides, run the checks, print every issue, save. Returns the issues."""
        # python-pptx relates the notes master it creates for speaker notes but does not list it in presentation.xml;
        # PowerPoint then asks to repair the file and may drop its pictures.
        pres = self.prs.part._element
        rid = next((r.rId for r in self.prs.part.rels.values() if r.reltype.endswith("/notesMaster")), None)
        if rid and pres.find(qn("p:notesMasterIdLst")) is None:
            pres.find(qn("p:sldMasterIdLst")).addnext(parse_xml(
                f'<p:notesMasterIdLst {nsdecls("p", "r")}><p:notesMasterId r:id="{rid}"/></p:notesMasterIdLst>'))
        for n, s in enumerate(self.prs.slides, 1):
            for t in s.shapes._spTree.iter("{http://schemas.openxmlformats.org/drawingml/2006/main}fld"):
                if t.get("type") == "slidenum":
                    t.find("{http://schemas.openxmlformats.org/drawingml/2006/main}t").text = str(n)
        self.prs.save(str(path))
        if not check:
            return []
        from .check import check as run
        issues = run(path)
        for i in issues:
            print(f"SLIDE CHECK {Path(path).name}: {i}")
        return issues

    # internals ------------------------------------------------------------------------------------------------

    def _slide(self, layout, title, notes):
        s = self.prs.slides.add_slide(self.layouts[layout])
        if title is not None:
            s.shapes.title.text_frame.text = title
        if layout in L.NUMBERED:   # python-pptx does not copy the slide-number placeholder from the layout
            num = next(p for p in self.layouts[layout].placeholders if p.placeholder_format.idx == 12)
            sp = copy.deepcopy(num._element)
            sp.spPr.clear()        # position and style come from the layout
            s.shapes._spTree.append(sp)
        if notes:
            s.notes_slide.notes_text_frame.text = notes
        return s

    def _place(self, slide, frame, path, poster=None):
        path = Path(path)
        ph = slide.placeholders[FIGURE_IDX[frame]]
        if path.suffix.lower() in (".mp4", ".mov", ".m4v"):
            x, y, w, h = ph.left, ph.top, ph.width, ph.height
            ph._element.getparent().remove(ph._element)
            if poster is None:
                sys.path.insert(0, str(SKILL.parent / "figures"))
                from figstyle.video import poster as grab
                poster = grab(path)
            slide.shapes.add_movie(str(path), x, y, w, h, poster_frame_image=str(poster), mime_type="video/mp4")
        else:
            frame = ph.left, ph.top, ph.width, ph.height
            pic = ph.insert_picture(str(path))   # crops to fill the frame if the figure was made at another size
            # The template's placeholders carry no geometry. PowerPoint draws a picture that inherits none as nothing,
            # so give each picture its frame and a rectangle explicitly.
            pic.left, pic.top, pic.width, pic.height = frame
            pic._element.spPr.append(parse_xml(f'<a:prstGeom {nsdecls("a")} prst="rect"><a:avLst/></a:prstGeom>'))

    @staticmethod
    def _items(ph, items, heading):
        from pptx.enum.dml import MSO_THEME_COLOR
        tf = ph.text_frame
        tf.paragraphs[0].text = heading
        tf.paragraphs[0].runs[0].font.color.theme_color = MSO_THEME_COLOR.TEXT_2   # the secondary ink

        def add(text, level):
            p = tf.add_paragraph()
            p.text, p.level = text, level

        for item in items:
            if isinstance(item, (list, tuple)):
                for sub in item:
                    add(sub, 1)
            else:
                add(item, 0)
